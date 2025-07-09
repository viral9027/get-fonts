import asyncio
import hashlib
import json
import logging
import os
import resource
import secrets
import time
from io import BytesIO
from urllib.parse import urlparse, urlunparse

import aiofiles
import aiohttp
import openpyxl
import pandas as pd
from aiohttp import ClientSession, ClientTimeout
from fastapi import FastAPI, File, UploadFile, Form, Request, Response, Depends
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from fontTools.ttLib import TTFont
from playwright.async_api import async_playwright, TimeoutError as PlaywrightTimeoutError, Error as PlaywrightError
from pydantic import BaseModel
import psutil

# Configure logging with reduced verbosity
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)
logging.getLogger("playwright").setLevel(logging.WARNING)

# Set conservative file descriptor limit
try:
    soft_limit, hard_limit = resource.getrlimit(resource.RLIMIT_NOFILE)
    new_soft_limit = min(2048, hard_limit)
    resource.setrlimit(resource.RLIMIT_NOFILE, (new_soft_limit, hard_limit))
    logger.info(f"File descriptor limit set to: {new_soft_limit}")
except Exception as e:
    logger.error(f"Failed to set file descriptor limit: {str(e)}")

app = FastAPI()
app.mount("/static", StaticFiles(directory="static"), name="static")
templates = Jinja2Templates(directory="templates")

# Session storage
users_db = {
    "user@example.com": {
        "email": "user@example.com",
        "hashed_password": hashlib.sha256("password123".encode()).hexdigest()
    }
}
sessions_db = {}
SESSION_FILE = "sessions.json"
FONT_DATA_FILE = "font_data.json"

async def load_sessions():
    try:
        async with aiofiles.open(SESSION_FILE, "r") as f:
            content = await f.read()
            return json.loads(content) if content else {}
    except Exception as e:
        logger.error(f"Error loading sessions: {str(e)}")
        return {}

async def save_sessions():
    try:
        async with aiofiles.open(SESSION_FILE, "w") as f:
            await f.write(json.dumps(sessions_db, indent=4))
    except Exception as e:
        logger.error(f"Error saving sessions: {str(e)}")

async def save_font_data(data_type: str, data: dict):
    try:
        existing_data = {"uploaded_fonts": [], "fetched_fonts": [], "bulk_fetched": []}
        if os.path.exists(FONT_DATA_FILE):
            async with aiofiles.open(FONT_DATA_FILE, "r") as f:
                content = await f.read()
                if content:
                    existing_data = json.loads(content)

        if data_type == "uploaded":
            existing_data["uploaded_fonts"].append(data)
        elif data_type == "fetched":
            existing_data["fetched_fonts"].append(data)
        elif data_type == "bulk_fetched":
            existing_data["bulk_fetched"] = data

        async with aiofiles.open(FONT_DATA_FILE, "w") as f:
            await f.write(json.dumps(existing_data, indent=4))
    except Exception as e:
        logger.error(f"Error saving font data: {str(e)}")

class LoginData(BaseModel):
    email: str
    password: str

def get_current_user(request: Request):
    session_id = request.cookies.get("session_id")
    if not session_id:
        logger.warning("No session_id cookie found")
        return None
    if session_id not in sessions_db:
        logger.warning(f"Session_id {session_id} not found")
        return None
    session_data = sessions_db[session_id]
    if time.time() - session_data["created_at"] > 86400:  # 24 hours
        logger.info(f"Session expired for session_id: {session_id}")
        del sessions_db[session_id]
        asyncio.create_task(save_sessions())
        return None
    logger.info(f"Valid session for user: {session_data['email']}")
    return session_data["email"]

def normalize_url(url: str) -> str:
    url = url.strip()
    if not url.startswith(('http://', 'https://')):
        url = 'https://' + url
    parsed = urlparse(url)
    if not parsed.netloc:
        raise ValueError("Invalid URL: No netloc specified")
    return urlunparse((
        parsed.scheme or 'https',
        parsed.netloc,
        parsed.path or '/',
        parsed.params,
        parsed.query,
        parsed.fragment
    ))

def extract_company_from_url(url: str) -> str:
    parsed = urlparse(url)
    domain = parsed.netloc.lower().replace("www.", "")
    return ' '.join(word.capitalize() for word in domain.split('.')[0].split('-'))

def extract_font_details(font: TTFont):
    name_table = font["name"]
    details = {}
    for record in name_table.names:
        if record.nameID in [0, 1, 2, 4, 6, 8, 9, 13]:
            try:
                details[record.nameID] = record.toUnicode()
            except:
                details[record.nameID] = record.string.decode("latin-1", errors="ignore")
    weight = str(font["OS/2"].usWeightClass) if "OS/2" in font else "Unknown"
    return {
        "family": details.get(1, "Unknown"),
        "subfamily": details.get(2, "Unknown"),
        "full_name": details.get(4, "Unknown"),
        "postscript_name": details.get(6, "Unknown"),
        "copyright": details.get(0, "Unknown"),
        "manufacturer": details.get(8, "Unknown"),
        "designer": details.get(9, "Unknown"),
        "license_type": details.get(13, "Unknown"),
        "weight": weight
    }

async def fetch_fonts_from_url(url: str, session: ClientSession, browser_context):
    font_details_list = []
    fonts = []
    cors_blocked = False

    page = await browser_context.new_page()

    async def capture_fonts(request):
        if request.resource_type == "font":
            font_url = request.url.lower()
            if any(font_url.endswith(ext) for ext in ['.ttf', '.otf', '.woff', '.woff2']):
                fonts.append(request.url)

    async def check_response(response):
        if response.request.resource_type == "font" and response.status in [403, 401]:
            nonlocal cors_blocked
            cors_blocked = True

    page.on("request", capture_fonts)
    page.on("response", check_response)
    await page.route("**/*", lambda route: route.abort() if route.request.resource_type in ["image", "media"] else route.continue_())

    try:
        # Log memory usage
        process = psutil.Process()
        mem = process.memory_info().rss / 1024 / 1024  # MB
        logger.info(f"Memory usage before navigating {url}: {mem:.2f}MB")

        await page.goto(url, wait_until="networkidle", timeout=30000)
        await page.wait_for_load_state("load", timeout=5000)
        await page.evaluate("document.fonts.ready")
    except (PlaywrightTimeoutError, PlaywrightError) as e:
        logger.warning(f"Navigation error for {url}: {str(e)}")
        if "Target closed" in str(e):
            logger.error(f"TargetClosedError for {url}, continuing with captured fonts")
    finally:
        await page.close()

    if not fonts:
        error_msg = "No downloadable fonts found" + (" (CORS restricted)" if cors_blocked else "")
        return [], error_msg

    for font_url in fonts:
        try:
            async with session.get(font_url, timeout=ClientTimeout(total=10)) as response:
                if response.status != 200:
                    continue
                content = await response.read()
                temp_file_path = f"temp_font_{secrets.token_hex(4)}{os.path.splitext(font_url)[1]}"
                async with aiofiles.open(temp_file_path, "wb") as f:
                    await f.write(content)
                try:
                    font = TTFont(temp_file_path)
                    font_details = extract_font_details(font)
                    font_details["url"] = font_url
                    font_details_list.append(font_details)
                finally:
                    if os.path.exists(temp_file_path):
                        os.remove(temp_file_path)
        except Exception as e:
            logger.error(f"Error downloading font {font_url}: {str(e)}")

    error_msg = "No fonts downloaded" if not font_details_list else None
    return font_details_list, error_msg

@app.get("/", response_class=HTMLResponse)
async def get_login(request: Request):
    return templates.TemplateResponse("login.html", {"request": request})

@app.post("/login", response_class=HTMLResponse)
async def login(response: Response, request: Request, email: str = Form(""), password: str = Form("")):
    user = users_db.get(email)
    hashed_password = hashlib.sha256(password.encode()).hexdigest()
    if not user or user["hashed_password"] != hashed_password:
        logger.warning(f"Login failed for email: {email}")
        return templates.TemplateResponse("login.html", {
            "request": request,
            "error": "Invalid credentials"
        })
    session_id = secrets.token_hex(16)
    sessions_db[session_id] = {"email": email, "created_at": time.time()}
    asyncio.create_task(save_sessions())
    response = RedirectResponse(url="/main", status_code=303)
    response.set_cookie(key="session_id", value=session_id, httponly=True, secure=False, samesite="Lax")
    logger.info(f"Login successful, session_id: {session_id} for user: {email}")
    return response

@app.get("/main", response_class=HTMLResponse)
async def get_main(request: Request, current_user: str = Depends(get_current_user)):
    if not current_user:
        logger.warning("No valid user session for /main")
        return RedirectResponse(url="/", status_code=303)
    return templates.TemplateResponse("main.html", {"request": request})

@app.get("/logout", response_class=RedirectResponse)
async def logout(request: Request, response: Response):
    session_id = request.cookies.get("session_id")
    if session_id in sessions_db:
        logger.info(f"Logging out session_id: {session_id}")
        del sessions_db[session_id]
        asyncio.create_task(save_sessions())
    response = RedirectResponse(url="/", status_code=303)
    response.delete_cookie("session_id")
    return response

@app.post("/upload-font", response_class=HTMLResponse)
async def upload_font(request: Request, file: UploadFile = File(...)):
    try:
        logger.info(f"Uploading font file: {file.filename}")
        if not file.filename.lower().endswith(('.ttf', '.otf', '.woff', '.woff2')):
            logger.warning(f"Unsupported font format: {file.filename}")
            return templates.TemplateResponse("main.html", {
                "request": request,
                "error": "Unsupported font format."
            })
        content = await file.read()
        temp_file_path = f"temp_{file.filename}_{secrets.token_hex(4)}"
        async with aiofiles.open(temp_file_path, "wb") as f:
            await f.write(content)
        try:
            font = TTFont(temp_file_path)
            font_details = extract_font_details(font)
            await save_font_data("uploaded", {"filename": file.filename, "font_details": font_details})
            logger.info(f"Successfully uploaded font: {file.filename}")
            return templates.TemplateResponse("main.html", {
                "request": request,
                "font_details": font_details,
                "filename": file.filename
            })
        finally:
            if os.path.exists(temp_file_path):
                os.remove(temp_file_path)
    except Exception as e:
        logger.error(f"Error processing font: {str(e)}")
        return templates.TemplateResponse("main.html", {
            "request": request,
            "error": f"Error processing font: {str(e)}"
        })

@app.post("/fetch-fonts", response_class=HTMLResponse)
async def fetch_fonts(request: Request, url: str = Form(...)):
    try:
        logger.info(f"Fetching fonts from {url}")
        normalized_url = normalize_url(url)
        async with aiohttp.ClientSession(timeout=ClientTimeout(total=10)) as session:
            async with async_playwright() as p:
                browser = await p.chromium.launch(
                    headless=True,
                    args=[
                        '--headless=new',
                        '--disable-gpu',
                        '--no-sandbox',
                        '--disable-dev-shm-usage',
                        '--disable-background-networking'
                    ]
                )
                try:
                    context = await browser.new_context(
                        viewport={'width': 1280, 'height': 720},
                        user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/126.0.0.0",
                        bypass_csp=True,
                        ignore_https_errors=True,
                        java_script_enabled=True
                    )
                    font_details_list, error_msg = await fetch_fonts_from_url(normalized_url, session, context)
                    await context.close()
                finally:
                    await browser.close()
        company = extract_company_from_url(normalized_url)
        result = {
            "website_url": normalized_url,
            "company": company,
            "total_fonts": len(font_details_list),
            "font_details_list": font_details_list,
            "error": error_msg
        }
        await save_font_data("fetched", result)
        logger.info(f"Fetched {len(font_details_list)} fonts from {normalized_url}")
        return templates.TemplateResponse("main.html", {
            "request": request,
            "bulk_results": [result],
            "message": f"Fetched {len(font_details_list)} fonts from {normalized_url}"
        })
    except Exception as e:
        logger.error(f"Error fetching fonts from {url}: {str(e)}")
        return templates.TemplateResponse("main.html", {
            "request": request,
            "error": f"Error fetching fonts: {str(e)}"
        })

@app.post("/upload-file", response_class=HTMLResponse)
async def upload_file(request: Request, file: UploadFile = File(...)):
    try:
        logger.info(f"Uploading file: {file.filename}")
        if not file.filename.endswith(('.csv', '.xlsx')):
            logger.warning(f"Invalid file format: {file.filename}")
            return templates.TemplateResponse("main.html", {
                "request": request,
                "error": "Please upload a CSV or XLSX file."
            })
        content = await file.read()
        temp_file_path = f"temp_{file.filename}_{secrets.token_hex(4)}"
        async with aiofiles.open(temp_file_path, "wb") as f:
            await f.write(content)
        try:
            df = pd.read_csv(temp_file_path) if file.filename.endswith('.csv') else pd.read_excel(temp_file_path, engine='openpyxl')
        finally:
            if os.path.exists(temp_file_path):
                os.remove(temp_file_path)
                logger.info(f"Removed temporary file: {temp_file_path}")

        company_col = next((col for col in df.columns if col.lower() in ["company", "name"]), None)
        website_col = next((col for col in df.columns if col.lower() in ["website", "url", "site"]), None)
        if not company_col or not website_col:
            logger.warning("Missing required columns in uploaded file")
            return templates.TemplateResponse("main.html", {
                "request": request,
                "error": "File must contain columns for company and website."
            })

        queue = [(row[company_col].strip(), row[website_col].strip()) for _, row in df.iterrows()]
        logger.info(f"Processing {len(queue)} URLs")

        async def process_url(company: str, website: str, session: ClientSession, browser_context, semaphore: asyncio.Semaphore):
            async with semaphore:
                try:
                    normalized_url = normalize_url(website)
                    font_details_list, error_msg = await asyncio.wait_for(
                        fetch_fonts_from_url(normalized_url, session, browser_context), timeout=30
                    )
                    logger.info(f"Processed {normalized_url}: {len(font_details_list)} fonts")
                    return {
                        "company": company,
                        "website_url": normalized_url,
                        "total_fonts": len(font_details_list),
                        "font_details_list": font_details_list,
                        "error": error_msg
                    }
                except Exception as e:
                    logger.error(f"Error processing {website}: {str(e)}")
                    return {
                        "company": company,
                        "website_url": website,
                        "total_fonts": 0,
                        "font_details_list": [],
                        "error": str(e)
                    }

        max_concurrent_tasks = 5
        bulk_results = []
        async with aiohttp.ClientSession(timeout=ClientTimeout(total=15), connector=aiohttp.TCPConnector(limit=50)) as session:
            async with async_playwright() as p:
                browser = await p.chromium.launch(
                    headless=True,
                    args=[
                        '--headless=new',
                        '--disable-gpu',
                        '--no-sandbox',
                        '--disable-dev-shm-usage',
                        '--disable-background-networking'
                    ]
                )
                try:
                    context = await browser.new_context(
                        viewport={'width': 1280, 'height': 720},
                        user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/126.0.0.0",
                        bypass_csp=True,
                        ignore_https_errors=True,
                        java_script_enabled=True
                    )
                    semaphore = asyncio.Semaphore(max_concurrent_tasks)
                    tasks = [process_url(company, website, session, context, semaphore) for company, website in queue]
                    bulk_results = await asyncio.gather(*tasks, return_exceptions=True)
                    bulk_results = [result for result in bulk_results if isinstance(result, dict)]
                    await context.close()
                finally:
                    await browser.close()
        await save_font_data("bulk_fetched", bulk_results)
        logger.info(f"Completed processing {len(bulk_results)} URLs")
        return templates.TemplateResponse("main.html", {
            "request": request,
            "bulk_results": bulk_results,
            "message": f"Processed {len(queue)} URLs"
        })
    except Exception as e:
        logger.error(f"Error processing file {file.filename}: {str(e)}")
        return templates.TemplateResponse("main.html", {
            "request": request,
            "error": f"Error processing file: {str(e)}"
        })

@app.get("/download-font-data", response_class=StreamingResponse)
async def download_font_data(request: Request):
    try:
        if not os.path.exists(FONT_DATA_FILE):
            logger.warning("No font data available to download")
            return templates.TemplateResponse("main.html", {
                "request": request,
                "error": "No font data available to download."
            })
        async with aiofiles.open(FONT_DATA_FILE, "r") as f:
            font_data = json.loads(await f.read())

        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        headers = [
            "Company", "Website URL", "Total Fonts", "Font Name", "Family",
            "Subfamily", "Weight", "Designer", "Manufacturer", "Copyright",
            "Font URL", "License Type", "Error"
        ]

        ws_uploaded = wb.create_sheet("Uploaded Fonts")
        ws_uploaded.append(headers)
        for font in font_data.get("uploaded_fonts", []):
            font_details = font.get("font_details", {})
            ws_uploaded.append([
                "N/A", "N/A", 1, font_details.get("full_name", "Unknown"),
                font_details.get("family", "Unknown"), font_details.get("subfamily", "Unknown"),
                font_details.get("weight", "Unknown"), font_details.get("designer", "Unknown"),
                font_details.get("manufacturer", "Unknown"), font_details.get("copyright", "Unknown"),
                "N/A", font_details.get("license_type", "Unknown"), "-"
            ])

        ws_fetched = wb.create_sheet("Fetched Fonts")
        ws_fetched.append(headers)
        for fetched_data in font_data.get("fetched_fonts", []):
            for font in fetched_data.get("font_details_list", []):
                ws_fetched.append([
                    fetched_data.get("company", "Unknown"), fetched_data.get("website_url", "Unknown"),
                    fetched_data.get("total_fonts", 0), font.get("full_name", "Unknown"),
                    font.get("family", "Unknown"), font.get("subfamily", "Unknown"),
                    font.get("weight", "Unknown"), font.get("designer", "Unknown"),
                    font.get("manufacturer", "Unknown"), font.get("copyright", "Unknown"),
                    font.get("url", "N/A"), font.get("license_type", "Unknown"), "-"
                ])

        ws_bulk = wb.create_sheet("Bulk Fetched Fonts")
        ws_bulk.append(headers)
        for result in font_data.get("bulk_fetched", []):
            if result.get("error"):
                ws_bulk.append([result.get("company", "Unknown"), result.get("website_url", "Unknown"),
                                0, "N/A", "N/A", "N/A", "N/A", "N/A", "N/A", "N/A", "N/A", "N/A", result.get("error", "-")])
            else:
                for font in result.get("font_details_list", []):
                    ws_bulk.append([
                        result.get("company", "Unknown"), result.get("website_url", "Unknown"),
                        result.get("total_fonts", 0), font.get("full_name", "Unknown"),
                        font.get("family", "Unknown"), font.get("subfamily", "Unknown"),
                        font.get("weight", "Unknown"), font.get("designer", "Unknown"),
                        font.get("manufacturer", "Unknown"), font.get("copyright", "Unknown"),
                        font.get("url", "N/A"), font.get("license_type", "Unknown"), "-"
                    ])

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        logger.info("Font data Excel file generated successfully")
        return StreamingResponse(
            content=output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=font_data.xlsx"}
        )
    except Exception as e:
        logger.error(f"Error generating Excel file: {str(e)}")
        return templates.TemplateResponse("main.html", {
            "request": request,
            "error": f"Error generating Excel file: {str(e)}"
        })

@app.post("/clear-font-data", response_class=HTMLResponse)
async def clear_font_data(request: Request):
    try:
        async with aiofiles.open(FONT_DATA_FILE, "w") as f:
            await f.write(json.dumps({"uploaded_fonts": [], "fetched_fonts": [], "bulk_fetched": []}, indent=4))
        logger.info("Font data cleared successfully")
        return templates.TemplateResponse("main.html", {
            "request": request,
            "message": "All font data cleared successfully."
        })
    except Exception as e:
        logger.error(f"Error clearing font data: {str(e)}")
        return templates.TemplateResponse("main.html", {
            "request": request,
            "error": f"Error clearing font data: {str(e)}"
        })

if __name__ == "__main__":
    import uvicorn
    asyncio.run(load_sessions())
    uvicorn.run("app:app", host="0.0.0.0", port=5000, workers=2)